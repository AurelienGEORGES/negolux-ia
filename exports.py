"""Extractions des réponses : Excel et CSV pour les tableaux, PDF pour la réponse complète.

Deux sources de tableaux :
- les tableaux Markdown écrits par Claude dans sa réponse ;
- les données brutes renvoyées par les outils MCP (« lignes », « detail », « totaux »...).
"""
from __future__ import annotations

import io
import json
import numbers
import re
from datetime import datetime
from pathlib import Path

import pandas as pd

MIME = {
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
}
CLES_TABLEAUX = ("lignes", "detail", "tables")
MAX_LIGNES_PDF = 500

# ---------- Lecture des tableaux ----------

_LIGNE = re.compile(r"^\s*\|.*\|\s*$")
_SEPARATEUR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")
_NOMBRE = re.compile(r"^[+-]?(\d{1,3}(?:[ \u00a0\u202f]\d{3})+|\d+)(?:[.,]\d+)?$")


def _cellules(ligne: str) -> list[str]:
    return [c.strip() for c in re.split(r"(?<!\\)\|", ligne.strip().strip("|"))]


def _valeur(cellule: str):
    """'1 234,5 €' -> 1234.5 ; les montants avec unité (k€, M€, %) restent du texte."""
    texte = cellule.replace("**", "").replace("`", "").replace("\\|", "|").strip()
    brut = texte.removesuffix("€").strip()
    if _NOMBRE.match(brut):
        nombre = re.sub(r"[ \u00a0\u202f]", "", brut).replace(",", ".")
        return int(nombre) if "." not in nombre else float(nombre)
    return texte


def _noms_uniques(noms: list[str]) -> list[str]:
    vus: dict[str, int] = {}
    resultat = []
    for nom in noms:
        nom = nom.replace("**", "") or "colonne"
        vus[nom] = vus.get(nom, 0) + 1
        resultat.append(nom if vus[nom] == 1 else f"{nom} ({vus[nom]})")
    return resultat


def tableaux_markdown(texte: str) -> list[pd.DataFrame]:
    """Tableaux Markdown présents dans une réponse."""
    lignes, tableaux, i = texte.splitlines(), [], 0
    while i < len(lignes) - 1:
        if _LIGNE.match(lignes[i]) and _SEPARATEUR.match(lignes[i + 1]):
            entete = _noms_uniques(_cellules(lignes[i]))
            i += 2
            corps = []
            while i < len(lignes) and _LIGNE.match(lignes[i]):
                cellules = _cellules(lignes[i])
                corps.append([_valeur(c) for c in (cellules + [""] * len(entete))[: len(entete)]])
                i += 1
            tableaux.append(pd.DataFrame(corps, columns=entete))
        else:
            i += 1
    return tableaux


def tableaux_outils(appels: list[dict]) -> list[tuple[str, pd.DataFrame]]:
    """Données brutes renvoyées par les outils MCP, sous forme de tableaux titrés."""
    tableaux = []
    for appel in appels:
        try:
            donnees = json.loads(appel.get("texte") or "")
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(donnees, dict):
            continue
        for cle in CLES_TABLEAUX:
            valeur = donnees.get(cle)
            if isinstance(valeur, list) and valeur and all(isinstance(v, dict) for v in valeur):
                tableaux.append((f"{appel['nom']} · {cle}", pd.json_normalize(valeur)))
        if isinstance(donnees.get("totaux"), dict):
            tableaux.append((f"{appel['nom']} · totaux", pd.json_normalize(donnees["totaux"])))
    return tableaux


def format_demande(question: str) -> str | None:
    """Format d'export que la question réclame, pour mettre le bon bouton en avant."""
    q = question.lower()
    if re.search(r"\b(excel|xlsx|classeur|tableur)\b", q):
        return "xlsx"
    if re.search(r"\bcsv\b", q):
        return "csv"
    if re.search(r"\b(pdf|rapport)\b", q):
        return "pdf"
    return None


# ---------- Excel et CSV ----------

def en_csv(df: pd.DataFrame) -> bytes:
    """CSV pour Excel en français : séparateur « ; », virgule décimale, UTF-8 avec BOM."""
    def virgule(v):  # cellule par cellule : les colonnes mêlent parfois nombres et texte
        return str(v).replace(".", ",") if isinstance(v, float) and not pd.isna(v) else v
    return df.map(virgule).to_csv(index=False, sep=";").encode("utf-8-sig")


def _nom_feuille(titre: str, pris: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", titre)[:31] or "Feuille"
    nom, n = base, 2
    while nom.lower() in pris:
        suffixe = f" ({n})"
        nom, n = base[: 31 - len(suffixe)] + suffixe, n + 1
    pris.add(nom.lower())
    return nom


def en_excel(tableaux: list[tuple[str, pd.DataFrame]]) -> bytes:
    """Un classeur, une feuille par tableau."""
    from openpyxl.styles import Font, PatternFill

    tampon, pris = io.BytesIO(), set()
    with pd.ExcelWriter(tampon, engine="openpyxl") as classeur:
        for titre, df in tableaux:
            nom = _nom_feuille(titre, pris)
            df.to_excel(classeur, sheet_name=nom, index=False)
            feuille = classeur.sheets[nom]
            feuille.freeze_panes = "A2"
            for cellule in feuille[1]:
                cellule.font = Font(bold=True, color="FFFFFF")
                cellule.fill = PatternFill("solid", fgColor="305496")
            for i, colonne in enumerate(df.columns, 1):
                largeur = max([len(str(colonne))] + [len(str(v)) for v in df[colonne].head(200)])
                feuille.column_dimensions[feuille.cell(1, i).column_letter].width = min(largeur + 2, 60)
    return tampon.getvalue()


# ---------- PDF ----------

POLICES = [  # (normal, gras, italique, gras italique) : Windows, Linux, macOS
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf",
     "C:/Windows/Fonts/arialbi.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-BoldOblique.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
     "/System/Library/Fonts/Supplemental/Arial Italic.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf"),
]
_HORS_BMP = re.compile(r"[\U00010000-\U0010FFFF\uFE0F\u200D]")  # émojis : absents des polices de texte


def _police(pdf) -> str:
    """Police Unicode du système si possible (accents, €), sinon Helvetica.

    Le normal et le gras sont requis ; à défaut d'italique, on réutilise le normal ou le gras.
    """
    for normal, gras, italique, gras_italique in POLICES:
        if Path(normal).exists() and Path(gras).exists():
            pdf.add_font("texte", "", normal)
            pdf.add_font("texte", "B", gras)
            pdf.add_font("texte", "I", italique if Path(italique).exists() else normal)
            pdf.add_font("texte", "BI", gras_italique if Path(gras_italique).exists() else gras)
            return "texte"
    return "helvetica"


SYMBOLES = {"⚠": "(!)", "✅": "[OK]", "✔": "[OK]", "✓": "[OK]", "❌": "[X]", "✗": "[X]", "⭐": "*"}


def _nettoyer(texte: str, police: str) -> str:
    for symbole, remplacement in SYMBOLES.items():  # absents des polices de texte
        texte = texte.replace(symbole, remplacement)
    texte = _HORS_BMP.sub("", texte).replace("--", "–")  # « -- » souligne en markdown fpdf2
    if police == "helvetica":  # police de base : Latin-1 seulement
        texte = texte.replace("€", "EUR").encode("latin-1", "replace").decode("latin-1")
    return texte


def _inline(texte: str) -> str:
    """Markdown en ligne -> syntaxe fpdf2 (**gras**, __italique__)."""
    texte = texte.replace("`", "")
    texte = re.sub(r"\*\*\*([^*\n]+)\*\*\*", r"**__\1__**", texte)
    texte = re.sub(r"(?<![*\w])\*(?!\*)([^*\n]+?)(?<!\*)\*(?![*\w])", r"__\1__", texte)
    return re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", texte)


def _texte_cellule(v) -> str:
    """Nombres à la française (1 234,56) ; accepte aussi les types numpy."""
    if v is None or isinstance(v, bool):
        return "" if v is None else str(v)
    if isinstance(v, numbers.Integral):
        return f"{int(v):,}".replace(",", " ")
    if isinstance(v, numbers.Real):
        return "" if pd.isna(v) else f"{float(v):,.2f}".replace(",", " ").replace(".", ",")
    return str(v)


def _tableau_pdf(pdf, police: str, df: pd.DataFrame) -> None:
    from fpdf.fonts import FontFace

    n = len(df.columns)
    pdf.set_font(police, size=9 if n <= 5 else 8 if n <= 8 else 7)
    numeriques = [  # colonne à droite si au moins la moitié des valeurs sont des nombres
        sum(isinstance(v, numbers.Number) and not isinstance(v, bool) for v in df[c]) * 2 >= len(df)
        for c in df.columns
    ]
    alignements = tuple("RIGHT" if num else "LEFT" for num in numeriques)
    entete = FontFace(emphasis="BOLD", color=255, fill_color=(48, 84, 150))
    with pdf.table(text_align=alignements, headings_style=entete, line_height=pdf.font_size * 1.6,
                   first_row_as_headings=True) as tableau:
        ligne = tableau.row()
        for colonne in df.columns:
            ligne.cell(_nettoyer(str(colonne), police))
        for valeurs in df.head(MAX_LIGNES_PDF).itertuples(index=False):
            ligne = tableau.row()
            for v in valeurs:
                ligne.cell(_nettoyer(_texte_cellule(v), police))
    if len(df) > MAX_LIGNES_PDF:
        pdf.set_font(police, "I", 8)
        pdf.multi_cell(0, 5, f"… {len(df) - MAX_LIGNES_PDF} lignes de plus dans l'export Excel.",
                       new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)


def en_pdf(question: str, reponse: str, utilisateur: str, date: datetime) -> bytes:
    """Rapport PDF : la question, puis la réponse avec ses tableaux."""
    from fpdf import FPDF

    class Rapport(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font(police, size=8)
            self.set_text_color(120)
            self.cell(0, 8, _nettoyer(f"Negolux · Assistant ERP — page {self.page_no()}/{{nb}}", police), align="C")

    pdf = Rapport(format="A4")
    police = _police(pdf)
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()

    pdf.set_font(police, "B", 16)
    pdf.multi_cell(0, 9, _nettoyer("Negolux · Réponse de l'assistant", police), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(police, size=9)
    pdf.set_text_color(110)
    pdf.multi_cell(0, 5, _nettoyer(f"Question posée le {date:%d/%m/%Y à %H:%M} par {utilisateur}", police),
                   new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    pdf.ln(3)

    pdf.set_font(police, "B", 12)
    pdf.multi_cell(0, 7, "Question", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(police, size=10)
    pdf.multi_cell(0, 5, _nettoyer(question or "(fichiers joints)", police), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_font(police, "B", 12)
    pdf.multi_cell(0, 7, "Réponse", new_x="LMARGIN", new_y="NEXT")

    lignes, i, paragraphe = reponse.splitlines(), 0, []

    def vider():
        if paragraphe:
            pdf.set_font(police, size=10)
            pdf.multi_cell(0, 5, _nettoyer(_inline("\n".join(paragraphe)), police), markdown=True,
                           new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            paragraphe.clear()

    while i < len(lignes):
        ligne = lignes[i]
        if _LIGNE.match(ligne) and i + 1 < len(lignes) and _SEPARATEUR.match(lignes[i + 1]):
            vider()
            fin = i + 2
            while fin < len(lignes) and _LIGNE.match(lignes[fin]):
                fin += 1
            _tableau_pdf(pdf, police, tableaux_markdown("\n".join(lignes[i:fin]))[0])
            i = fin
            continue
        titre = re.match(r"^(#{1,6})\s+(.*)", ligne)
        if titre:
            vider()
            pdf.set_font(police, "B", 13 if len(titre.group(1)) <= 2 else 11)
            pdf.multi_cell(0, 6, _nettoyer(titre.group(2).replace("**", ""), police), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        elif not ligne.strip() or re.match(r"^\s*([-*_])\s*(\1\s*){2,}$", ligne):  # vide ou ligne « --- »
            vider()
        else:
            puce = re.match(r"^(\s*)[-*]\s+(.*)", ligne)
            paragraphe.append(f"{puce.group(1)}•  {puce.group(2)}" if puce else ligne)
        i += 1
    vider()
    return bytes(pdf.output())


# ---------- Tout préparer pour une réponse ----------

def preparer(question: str, reponse: str, appels: list[dict], utilisateur: str, date: datetime) -> dict:
    """Fichiers proposés au téléchargement : le PDF toujours, Excel et CSV s'il y a des tableaux."""
    tableaux = [(f"Tableau {n}", df) for n, df in enumerate(tableaux_markdown(reponse), 1)]
    tableaux += tableaux_outils(appels)
    base = f"negolux_{date:%Y%m%d_%H%M%S}"
    return {
        "base": base,
        "pdf": en_pdf(question, reponse, utilisateur, date),
        "xlsx": en_excel(tableaux) if tableaux else None,
        "csv": [(titre, en_csv(df)) for titre, df in tableaux],
    }


def nom_csv(base: str, titre: str) -> str:
    return f"{base}_{re.sub(r'[^a-z0-9]+', '_', titre.lower()).strip('_')}.csv"
