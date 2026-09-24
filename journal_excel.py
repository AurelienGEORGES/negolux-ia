"""Journal des échanges avec Claude, commun à l'interface et au terminal.

- journal/echanges.jsonl : la source de vérité, une ligne JSON par échange ;
- journal/echanges.xlsx : le classeur, régénéré à partir du JSONL après chaque ajout.

Si le classeur est ouvert dans Excel, il ne peut pas être réécrit : l'échange reste dans
echanges.jsonl et le classeur est rattrapé au prochain ajout, ou à la main avec
    python journal_excel.py
"""
from __future__ import annotations

import json
import os
import time
from collections import Counter
from contextlib import contextmanager, suppress
from datetime import datetime
from pathlib import Path
from typing import Iterable

RACINE = Path(__file__).resolve().parent
JOURNAL = RACINE / "journal"
JSONL = JOURNAL / "echanges.jsonl"
XLSX = JOURNAL / "echanges.xlsx"
VERROU = JOURNAL / ".verrou"

MAX_CELLULE = 32_767  # limite d'Excel par cellule
# (clé dans le JSONL, titre, largeur) ; N° et Durée sont calculés
COLONNES = [
    ("n", "N°", 6), ("date_question", "Date question", 19), ("date_reponse", "Date réponse", 19),
    ("duree", "Durée (s)", 10), ("source", "Source", 11), ("utilisateur", "Utilisateur", 14),
    ("session", "Session", 10), ("question", "Question", 60), ("pieces_jointes", "Pièces jointes", 30),
    ("reponse", "Réponse", 100), ("outils", "Outils utilisés", 40), ("statut", "Statut", 12),
]


def maintenant() -> str:
    return datetime.now().isoformat(timespec="seconds")


def resume_outils(noms: Iterable[str]) -> str:
    """['a', 'b', 'a'] -> 'a ×2, b'."""
    return ", ".join(f"{nom} ×{n}" if n > 1 else nom for nom, n in Counter(noms).items())


@contextmanager
def verrou():
    """Empêche deux écritures simultanées (interface et terminal, ou deux sessions)."""
    JOURNAL.mkdir(exist_ok=True)
    limite = time.time() + 10
    while True:
        try:
            fd = os.open(VERROU, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            break
        except FileExistsError:
            with suppress(FileNotFoundError):
                if time.time() - VERROU.stat().st_mtime > 60:  # verrou abandonné
                    VERROU.unlink()
                    continue
            if time.time() > limite:
                raise TimeoutError("journal verrouillé depuis plus de 10 s")
            time.sleep(0.1)
    try:
        yield
    finally:
        os.close(fd)
        VERROU.unlink(missing_ok=True)


def lire() -> list[dict]:
    """Tous les échanges, du plus ancien au plus récent."""
    echanges = []
    if JSONL.exists():
        for ligne in JSONL.read_text(encoding="utf-8").splitlines():
            with suppress(json.JSONDecodeError):
                echanges.append(json.loads(ligne))
    return echanges


def ajouter(echanges: list[dict]) -> str | None:
    """Ajoute des échanges au journal ; renvoie un message si le classeur n'a pas pu être mis à jour."""
    with verrou():
        with open(JSONL, "a", encoding="utf-8") as f:
            for e in echanges:
                f.write(json.dumps(e, ensure_ascii=False) + "\n")
        return _reconstruire()


def reconstruire() -> str | None:
    with verrou():
        return _reconstruire()


def _reconstruire() -> str | None:
    """Réécrit echanges.xlsx à partir de echanges.jsonl ; renvoie un message si c'est impossible."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return ("Journal : openpyxl n'est pas installé (python -m pip install -r requirements.txt). "
                "L'échange est gardé dans journal/echanges.jsonl.")

    echanges = lire()

    def texte(valeur):
        valeur = ILLEGAL_CHARACTERS_RE.sub("", valeur or "")
        if len(valeur) > MAX_CELLULE:
            note = "\n[… tronqué : texte complet dans journal/echanges.jsonl]"
            valeur = valeur[: MAX_CELLULE - len(note)] + note
        return valeur

    def date(valeur):
        return datetime.fromisoformat(valeur) if valeur else None

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Échanges")
    for i, (_, _, largeur) in enumerate(COLONNES, 1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{len(echanges) + 1}"

    entete = []
    for _, titre, _ in COLONNES:
        c = WriteOnlyCell(ws, value=titre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        entete.append(c)
    ws.append(entete)

    haut = Alignment(vertical="top", wrap_text=True)
    for n, e in enumerate(echanges, 1):
        dq, dr = date(e.get("date_question")), date(e.get("date_reponse"))
        valeurs = {
            "n": n, "date_question": dq, "date_reponse": dr,
            "duree": round((dr - dq).total_seconds()) if dq and dr else None,
            "source": e.get("source", "terminal"),  # les premiers échanges venaient tous du terminal
            "session": (e.get("session") or "")[:8],
        }
        ligne = []
        for cle, _, _ in COLONNES:
            v = valeurs[cle] if cle in valeurs else e.get(cle)
            if cle in ("question", "reponse", "pieces_jointes", "outils"):
                v = texte(v)
            c = WriteOnlyCell(ws, value=v)
            c.alignment = haut
            if isinstance(v, datetime):
                c.number_format = "dd/mm/yyyy hh:mm:ss"
            ligne.append(c)
        ws.append(ligne)

    temporaire = XLSX.with_name("echanges.tmp.xlsx")
    wb.save(temporaire)
    try:
        os.replace(temporaire, XLSX)
    except OSError:  # classeur ouvert dans Excel (verrouillé sous Windows)
        temporaire.unlink(missing_ok=True)
        return "Journal : echanges.xlsx est ouvert dans Excel, il sera mis à jour au prochain échange."
    return None


if __name__ == "__main__":
    print(reconstruire() or f"{XLSX} reconstruit.")
