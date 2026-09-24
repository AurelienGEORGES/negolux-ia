"""Journal des échanges Claude Code : chaque question et chaque réponse dans journal/echanges.xlsx.

Appelé par les hooks de .claude/settings.json :
- UserPromptSubmit : met la question de côté (journal/.en_attente/<session>.json) ;
- Stop : retrouve la réponse dans le transcript de la session, ajoute l'échange à
  journal/echanges.jsonl (la source de vérité), puis régénère journal/echanges.xlsx.

Si le classeur est ouvert dans Excel, il ne peut pas être réécrit : l'échange reste dans
echanges.jsonl et le classeur est rattrapé à la réponse suivante, ou à la main avec
    python .claude/hooks/journal_echanges.py --reconstruire

Le script ne bloque jamais Claude Code : une erreur est notée dans journal/erreurs.log.
"""
import getpass
import json
import os
import sys
import time
import traceback
from collections import Counter
from contextlib import contextmanager, suppress
from datetime import datetime
from pathlib import Path

RACINE = Path(__file__).resolve().parents[2]
JOURNAL = RACINE / "journal"
EN_ATTENTE = JOURNAL / ".en_attente"
JSONL = JOURNAL / "echanges.jsonl"
XLSX = JOURNAL / "echanges.xlsx"
VERROU = JOURNAL / ".verrou"

MAX_CELLULE = 32_767        # limite d'Excel par cellule
PERIME_APRES_S = 6 * 3600   # question d'une autre session restée sans Stop depuis 6 h
COLONNES = [  # (titre, largeur)
    ("N°", 6), ("Date question", 19), ("Date réponse", 19), ("Durée (s)", 10),
    ("Utilisateur", 14), ("Session", 10), ("Question", 60), ("Réponse", 100),
    ("Outils utilisés", 40), ("Statut", 12),
]


def maintenant():
    return datetime.now().isoformat(timespec="seconds")


def heure_locale(horodatage_utc):
    """'2026-09-24T07:55:31.434Z' -> '2026-09-24T09:55:31' (heure locale)."""
    d = datetime.fromisoformat(horodatage_utc.replace("Z", "+00:00"))
    return d.astimezone().replace(tzinfo=None).isoformat(timespec="seconds")


@contextmanager
def verrou():
    """Empêche deux sessions d'écrire le journal en même temps."""
    limite = time.time() + 10
    while True:
        try:
            fd = os.open(VERROU, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            break
        except FileExistsError:
            with suppress(FileNotFoundError):
                if time.time() - VERROU.stat().st_mtime > 60:  # verrou abandonné
                    VERROU.unlink()
                    continue
            if time.time() > limite:
                raise TimeoutError("journal verrouillé depuis plus de 10 s")
            time.sleep(0.1)
    try:
        yield
    finally:
        os.close(fd)
        VERROU.unlink(missing_ok=True)


# ---------- Lecture du transcript ----------

def charger_transcript(chemin):
    if not chemin or not Path(chemin).exists():
        return []
    lignes = []
    with open(chemin, encoding="utf-8") as f:
        for ligne in f:
            with suppress(json.JSONDecodeError):
                lignes.append(json.loads(ligne))
    return lignes


def texte_utilisateur(ligne):
    """Texte d'une vraie question de l'utilisateur ; None pour toute autre ligne."""
    if ligne.get("type") != "user" or ligne.get("isMeta") or ligne.get("isSidechain"):
        return None
    contenu = ligne.get("message", {}).get("content")
    if isinstance(contenu, str):
        return contenu
    if isinstance(contenu, list) and not any(b.get("type") == "tool_result" for b in contenu):
        return "\n".join(b.get("text", "") for b in contenu if b.get("type") == "text")
    return None


def lire_reponse(transcript, questions, essais=1):
    """Réponse de Claude aux questions en attente : (texte, outils appelés, horodatage de fin).

    Tout ce que Claude a écrit après la première question en attente en fait partie,
    y compris après un message envoyé pendant qu'il travaillait.
    """
    for essai in range(essais):
        lignes = charger_transcript(transcript)
        posees = [(i, t) for i, l in enumerate(lignes) if (t := texte_utilisateur(l)) is not None]
        # Début : la première question en attente (à défaut, la dernière posée)
        debut = next((i for i, t in reversed(posees) if t.strip() == questions[0].strip()),
                     posees[-1][0] if posees else len(lignes))
        textes, outils, fin = [], Counter(), None
        for ligne in lignes[debut + 1:]:
            if ligne.get("type") != "assistant" or ligne.get("isSidechain"):
                continue
            contenu = ligne.get("message", {}).get("content") or []
            if isinstance(contenu, str):
                contenu = [{"type": "text", "text": contenu}]
            for bloc in contenu:
                if bloc.get("type") == "text" and bloc.get("text", "").strip():
                    textes.append(bloc["text"].strip())
                elif bloc.get("type") == "tool_use":
                    outils[bloc.get("name", "?")] += 1
            fin = ligne.get("timestamp") or fin
        if textes or essai == essais - 1:
            break
        time.sleep(0.3)  # le transcript n'est peut-être pas encore écrit
    liste_outils = ", ".join(f"{nom} ×{n}" if n > 1 else nom for nom, n in outils.items())
    return "\n\n".join(textes), liste_outils, fin


# ---------- Événements ----------

def fichier_attente(session):
    return EN_ATTENTE / f"{session}.json"


def mettre_de_cote(entree):
    """UserPromptSubmit : ajoute la question à la liste d'attente de la session."""
    f = fichier_attente(entree.get("session_id", "inconnue"))
    attente = json.loads(f.read_text(encoding="utf-8")) if f.exists() else []
    attente.append({
        "question": entree.get("prompt", ""),
        "date": maintenant(),
        "transcript": entree.get("transcript_path"),
    })
    f.write_text(json.dumps(attente, ensure_ascii=False), encoding="utf-8")


def echange(f, session, en_direct, dernier_message=None):
    """Construit l'échange à partir d'un fichier d'attente et du transcript."""
    attente = json.loads(f.read_text(encoding="utf-8"))
    questions = [a["question"] for a in attente]
    reponse, outils, fin = lire_reponse(attente[-1].get("transcript"), questions, essais=5 if en_direct else 1)
    if isinstance(dernier_message, str) and dernier_message.strip() and not reponse.endswith(dernier_message.strip()):
        reponse = f"{reponse}\n\n{dernier_message.strip()}".strip()
    if not reponse:
        statut = "sans réponse"
    else:
        statut = "répondu" if en_direct else "interrompu"
    return {
        "date_question": attente[0]["date"],
        "date_reponse": heure_locale(fin) if fin else (maintenant() if en_direct else ""),
        "utilisateur": getpass.getuser(),
        "session": session,
        "question": "\n\n— puis —\n\n".join(questions),
        "reponse": reponse,
        "outils": outils,
        "statut": statut,
    }


def enregistrer(entree):
    """Stop : enregistre l'échange de la session, plus ceux des sessions fermées sans Stop."""
    session = entree.get("session_id", "inconnue")
    a_traiter = []
    f = fichier_attente(session)
    if f.exists():
        a_traiter.append((f, echange(f, session, True, entree.get("last_assistant_message"))))
    for autre in EN_ATTENTE.glob("*.json"):
        if autre != f and time.time() - autre.stat().st_mtime > PERIME_APRES_S:
            a_traiter.append((autre, echange(autre, autre.stem, False)))
    if not a_traiter:
        return None
    with verrou():
        with open(JSONL, "a", encoding="utf-8") as j:
            for _, e in a_traiter:
                j.write(json.dumps(e, ensure_ascii=False) + "\n")
        for fichier, _ in a_traiter:
            fichier.unlink(missing_ok=True)
        return reconstruire_excel()


# ---------- Classeur Excel ----------

def reconstruire_excel():
    """Réécrit echanges.xlsx à partir de echanges.jsonl ; renvoie un message si c'est impossible."""
    try:
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return ("Journal des échanges : openpyxl n'est pas installé (pip install -r requirements.txt). "
                "L'échange est gardé dans journal/echanges.jsonl.")

    echanges = []
    if JSONL.exists():
        for ligne in JSONL.read_text(encoding="utf-8").splitlines():
            with suppress(json.JSONDecodeError):
                echanges.append(json.loads(ligne))

    def texte(valeur):
        valeur = ILLEGAL_CHARACTERS_RE.sub("", valeur or "")
        if len(valeur) > MAX_CELLULE:
            note = "\n[… tronqué : texte complet dans journal/echanges.jsonl]"
            valeur = valeur[: MAX_CELLULE - len(note)] + note
        return valeur

    def date(valeur):
        return datetime.fromisoformat(valeur) if valeur else None

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Échanges")
    for i, (_, largeur) in enumerate(COLONNES, 1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLONNES))}{len(echanges) + 1}"

    entete = []
    for titre, _ in COLONNES:
        c = WriteOnlyCell(ws, value=titre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        entete.append(c)
    ws.append(entete)

    haut = Alignment(vertical="top", wrap_text=True)
    for n, e in enumerate(echanges, 1):
        dq, dr = date(e.get("date_question")), date(e.get("date_reponse"))
        valeurs = [
            n, dq, dr, round((dr - dq).total_seconds()) if dq and dr else None,
            e.get("utilisateur"), (e.get("session") or "")[:8], texte(e.get("question")),
            texte(e.get("reponse")), e.get("outils"), e.get("statut"),
        ]
        ligne = []
        for v in valeurs:
            c = WriteOnlyCell(ws, value=v)
            c.alignment = haut
            if isinstance(v, datetime):
                c.number_format = "dd/mm/yyyy hh:mm:ss"
            ligne.append(c)
        ws.append(ligne)

    temporaire = XLSX.with_name("echanges.tmp.xlsx")
    wb.save(temporaire)
    try:
        os.replace(temporaire, XLSX)
    except OSError:  # classeur ouvert dans Excel (verrouillé sous Windows)
        temporaire.unlink(missing_ok=True)
        return "Journal des échanges : echanges.xlsx est ouvert, il sera mis à jour à la prochaine réponse."
    return None


def main():
    JOURNAL.mkdir(exist_ok=True)
    EN_ATTENTE.mkdir(exist_ok=True)
    if "--reconstruire" in sys.argv:
        with verrou():
            message = reconstruire_excel()
        print(message or f"{XLSX} reconstruit.")
        return
    brut = sys.stdin.buffer.read().decode("utf-8")  # UTF-8 aussi sous Windows
    entree = json.loads(brut) if brut.strip() else {}
    message = None
    if entree.get("hook_event_name") == "UserPromptSubmit":
        mettre_de_cote(entree)
    elif entree.get("hook_event_name") == "Stop":
        message = enregistrer(entree)
    if message:
        print(json.dumps({"systemMessage": message}))


if __name__ == "__main__":
    try:
        main()
    except Exception:
        JOURNAL.mkdir(exist_ok=True)
        with open(JOURNAL / "erreurs.log", "a", encoding="utf-8") as log:
            log.write(f"{maintenant()}\n{traceback.format_exc()}\n")
        print(json.dumps({"systemMessage": "Journal des échanges : erreur, voir journal/erreurs.log"}))
    sys.exit(0)
